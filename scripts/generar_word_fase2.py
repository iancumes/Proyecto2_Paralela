"""Renderiza el contenido del informe en Word y conserva la fuente Markdown."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from contenido_fase2 import generate

ROOT = Path(__file__).resolve().parents[1]


def build():
    markdown = generate(ROOT)
    (ROOT/'fase2/INFORME_FASE2.md').write_text(markdown, encoding='utf-8')
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2.0), Cm(2.0)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(7)
    doc.styles['Title'].font.size = Pt(23)
    for name, size in [('Heading 1',15), ('Heading 2',11.5)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Proyecto 2 fase 2  |  ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = 'Informe de la fase 2 del Proyecto 2'
    doc.core_properties.subject = 'Validacion, Docker, ensayo virtual y preparacion del cluster fisico'
    doc.core_properties.author = ''

    def add_table(lines):
        values = [[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
        rows = [values[0]] + values[2:]
        count = len(rows[0])
        widths = {2:[6.1,10.9], 4:[4.0,4.5,4.5,4.0], 5:[7.2,1.0,3.0,2.7,3.1]}.get(count, [17/count]*count)
        table = doc.add_table(rows=0, cols=count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        for i, row in enumerate(rows):
            cells = table.add_row().cells
            props = table.rows[-1]._tr.get_or_add_trPr()
            props.append(OxmlElement('w:cantSplit'))
            if i == 0:
                props.append(OxmlElement('w:tblHeader'))
            for j, (cell, text) in enumerate(zip(cells, row)):
                cell.text = text
                cell.width = Cm(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'DCE6F1' if i==0 else ('F5F7FA' if i%2==0 else 'FFFFFF'))
                tcpr.append(shade)
                borders = OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    border = OxmlElement('w:'+edge)
                    for key, value in [('val','single'),('sz','4'),('color','D9D9D9')]:
                        border.set(qn('w:'+key), value)
                    borders.append(border)
                tcpr.append(borders)
                margins = OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    m = OxmlElement('w:'+edge)
                    m.set(qn('w:w'), '85'); m.set(qn('w:type'), 'dxa')
                    margins.append(m)
                tcpr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.04
                    for run in p.runs:
                        run.font.size = Pt(9.2)
                        run.font.bold = i==0
        doc.add_paragraph().paragraph_format.space_after = Pt(1)

    lines = markdown.splitlines()
    index = 0
    in_code = False
    while index < len(lines):
        line = lines[index]
        if line.startswith('```'):
            in_code = not in_code
        elif in_code:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(line)
            r.font.name, r.font.size = 'Consolas', Pt(8.5)
        elif line.startswith('|'):
            group = []
            while index < len(lines) and lines[index].startswith('|'):
                group.append(lines[index]); index += 1
            add_table(group)
            continue
        elif line == '<!-- pagebreak -->':
            doc.add_page_break()
        elif line.startswith('# '):
            doc.add_paragraph(line[2:], style='Title')
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.strip():
            doc.add_paragraph(line)
        index += 1
    out = ROOT/'output/docx'
    out.mkdir(parents=True, exist_ok=True)
    target = out/'Informe_Proyecto2_Fase2.docx'
    doc.save(target)
    print(target)


if __name__ == '__main__':
    build()
